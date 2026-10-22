from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.visualizacion import tasa_cancelacion

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(22, 45, 76)
BLUE = RGBColor(49, 108, 170)
ORANGE = RGBColor(230, 126, 34)
GRAY = RGBColor(85, 96, 108)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.65))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.0), Inches(12.0), Inches(0.4))
        p = sub_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = GRAY


def add_bullets(
    slide,
    lines: list[str],
    top: float = 1.5,
    left: float = 0.85,
    width: float = 11.7,
    font_size: int = 18,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.4))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = GRAY
        paragraph.space_after = Pt(14)


def build_presentation(
    data_path: str | Path = ROOT / 'data' / 'processed' / 'netflix_limpio.csv',
    metrics_path: str | Path = ROOT / 'reports' / 'metrics' / 'metricas_modelos.csv',
    evaluation_path: str | Path = ROOT / 'reports' / 'metrics' / 'evaluation_netflix.json',
    importance_path: str | Path = ROOT / 'reports' / 'metrics' / 'variables_importantes.csv',
    output_path: str | Path = ROOT / 'reports' / 'presentacion_netflix.pptx',
) -> Path:
    data = pd.read_csv(data_path)
    metrics = pd.read_csv(metrics_path)
    evaluation = json.loads(Path(evaluation_path).read_text(encoding='utf-8'))
    importance = pd.read_csv(importance_path)
    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.8), Inches(1.3))
    p = title.text_frame.paragraphs[0]
    p.text = '¿Quién va a cancelar Netflix?'
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = slide.shapes.add_textbox(Inches(0.85), Inches(3.15), Inches(11.5), Inches(0.8))
    p = subtitle.text_frame.paragraphs[0]
    p.text = 'Análisis exploratorio y modelos de clasificación · Caso académico'
    p.font.size = Pt(19)
    p.font.color.rgb = RGBColor(225, 233, 242)
    note = slide.shapes.add_textbox(Inches(0.85), Inches(5.9), Inches(11.5), Inches(0.5))
    p = note.text_frame.paragraphs[0]
    p.text = f'Fuente: neflix.csv · {len(data):,} filas limpias · Métricas calculadas sobre prueba reservada'
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(225, 233, 242)

    target = data['cancelo'].dropna()
    canceled = int(target.sum())
    total = int(target.size)
    missing = data.isna().sum().sort_values(ascending=False)
    missing = missing[missing > 0]
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, 'Datos y alcance', 'Registros limpios exportados; los faltantes de predictores se imputan dentro del pipeline.')
    lines = [
        f'Archivo analizado: neflix.csv. Filas originales: 5,025; filas conservadas después de resolver duplicados: {len(data):,}.',
        f'Cancelaciones válidas: {canceled:,} de {total:,} usuarios ({canceled / total:.1%}).',
        f'Filas con valores faltantes en al menos una columna: {int(data.isna().any(axis=1).sum()):,}.',
        'La variable objetivo no se imputa; fecha_registro e id_usuario se excluyen como predictores.',
        'La asociación encontrada es descriptiva/predictiva y no demuestra causalidad ni validez en operación real.',
    ]
    if not missing.empty:
        lines.insert(3, 'Faltantes por variable: ' + ', '.join(f'{name}={int(value)}' for name, value in missing.items()))
    add_bullets(slide, lines)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, 'Perfiles descriptivos', 'Tasa y tamaño de cada grupo; comparar tasas pequeñas sin ignorar el número de usuarios.')
    groups = []
    for variable, label in [('plan', 'Plan'), ('pais', 'País'), ('metodo_pago', 'Método de pago'), ('dispositivo_principal', 'Dispositivo')]:
        table = tasa_cancelacion(data, variable).sort_values('tasa_cancelacion', ascending=False)
        row = table.iloc[0]
        groups.append(
            f'{label}: mayor tasa observada en {row[variable]} — {row["tasa_cancelacion"]:.1%} '
            f'({int(row["cancelados"])} de {int(row["usuarios"])} usuarios).'
        )
    medians = data.groupby('cancelo')[[
        'horas_vistas_mes', 'dias_desde_ultimo_acceso',
        'calificacion_promedio', 'tickets_soporte',
    ]].median()
    for name, column in [
        ('Horas vistas/mes', 'horas_vistas_mes'),
        ('Días desde último acceso', 'dias_desde_ultimo_acceso'),
        ('Calificación promedio', 'calificacion_promedio'),
        ('Tickets de soporte', 'tickets_soporte'),
    ]:
        no_cancel = medians.loc[0, column]
        yes_cancel = medians.loc[1, column]
        groups.append(f'{name}: mediana No cancela={no_cancel:g}; Sí cancela={yes_cancel:g}.')
    add_bullets(slide, groups, top=1.3, font_size=14)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, 'Evaluación de modelos', 'Prueba reservada 25%; los modelos se comparan usando cancelo=1 como clase positiva.')
    rows = [['Modelo', 'Exactitud', 'Precisión', 'Recall', 'F1', 'F1 CV train']]
    for _, row in metrics.iterrows():
        rows.append([
            str(row['modelo']),
            f'{row["accuracy"]:.3f}',
            f'{row["precision"]:.3f}',
            f'{row["recall"]:.3f}',
            f'{row["f1"]:.3f}',
            f'{row["train_cv_f1_mean"]:.3f}',
        ])
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.55), Inches(1.55), Inches(12.2), Inches(0.55 * len(rows)))
    table = table_shape.table
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.size = Pt(12 if r else 13)
                paragraph.font.bold = r == 0
                paragraph.font.color.rgb = RGBColor(255, 255, 255) if r == 0 else GRAY
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BLUE
    counts = evaluation['test_class_counts']
    note_text = (
        f'Prueba: {evaluation["test_rows"]:,} usuarios; clase 0={counts.get("0", 0):,}, '
        f'clase 1={counts.get("1", 0):,}. La exactitud puede ocultar un recall bajo en la clase minoritaria.'
    )
    note = slide.shapes.add_textbox(Inches(0.75), Inches(4.1), Inches(7.8), Inches(1.2))
    p = note.text_frame.paragraphs[0]
    p.text = note_text
    p.font.size = Pt(16)
    p.font.color.rgb = GRAY
    selected_model = str(evaluation['selected_for_permutation_importance']).lower()
    matrix = ROOT / 'reports' / 'figures' / f'matriz_confusion_{selected_model}.png'
    if matrix.exists():
        slide.shapes.add_picture(str(matrix), Inches(9.0), Inches(4.45), width=Inches(3.5))

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, 'Variables asociadas a la predicción', 'Importancia por permutación, métrica F1, calculada solo con entrenamiento.')
    chart = ROOT / 'reports' / 'figures' / 'variables_importantes.png'
    if chart.exists():
        slide.shapes.add_picture(str(chart), Inches(0.7), Inches(1.3), width=Inches(7.6))
    top = importance.head(5)
    lines = [f'{row.variable}: importancia media {row.importancia_media:.3f}' for row in top.itertuples()]
    lines.append('La importancia predictiva no implica que una variable cause la cancelación.')
    add_bullets(slide, lines, top=1.65, left=8.5, width=4.1, font_size=14)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, 'Acciones a validar', 'Propuestas de experimentación; no se atribuyen ahorros ni eficacia no medida.')
    lines = [
        'Probar una intervención de reactivación para usuarios con menor actividad reciente; medir retención a 30 días y tasa de cancelación del grupo.',
        'Revisar la experiencia de soporte en segmentos con más tickets; medir resolución en primer contacto y cancelación posterior.',
        'Auditar el flujo de pago en los grupos con mayor tasa descriptiva; medir fallos de pago y cancelación, controlando el tamaño de cada grupo.',
        'Validar cualquier campaña con grupo de control y definir la ventana temporal antes de atribuir resultados.',
        'Limitación temporal: el dataset no prueba que todas las variables estuvieran disponibles antes de cancelar.',
    ]
    add_bullets(slide, lines)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path


def main() -> None:
    print(f'Presentación guardada en {build_presentation()}')


if __name__ == '__main__':
    main()
